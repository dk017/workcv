"""Build public/downloads/job-application-tracker-uk.xlsx.

Columns match the online tracker's CSV export (lib/job-tracker.ts), so a sheet saved
as CSV can be imported into https://workcv.co.uk/tools/job-application-tracker-uk.
Run: python scripts/generate-job-tracker-template.py (needs openpyxl).
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "public" / "downloads" / "job-application-tracker-uk.xlsx"

COLUMNS = [
    ("Job title", 28), ("Employer", 24), ("Status", 14), ("Date applied", 14), ("Closing date", 14),
    ("Next action", 30), ("Next action date", 16), ("Location", 16), ("Salary", 14), ("Where found", 20),
    ("Job advert link", 32), ("Contact", 20), ("Notes", 36), ("Job advert text", 40),
]
STATUSES = ["Saved", "Applied", "Interview", "Offer", "Unsuccessful", "Withdrawn"]
SOURCES = ["Indeed", "Reed", "Totaljobs", "CV-Library", "LinkedIn", "NHS Jobs", "Civil Service Jobs",
           "Find a job (GOV.UK)", "Company website", "Recruitment agency", "Referral", "Other"]
ROWS = 300
NAVY = "0F2942"

wb = Workbook()
ws = wb.active
ws.title = "Applications"
for index, (header, width) in enumerate(COLUMNS, start=1):
    cell = ws.cell(row=1, column=index, value=header)
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.alignment = Alignment(vertical="center")
    ws.column_dimensions[cell.column_letter].width = width
ws.row_dimensions[1].height = 22
ws.freeze_panes = "C2"
ws.auto_filter.ref = f"A1:N{ROWS + 1}"

status_list = DataValidation(type="list", formula1=f'"{",".join(STATUSES)}"', allow_blank=True)
source_list = DataValidation(type="list", formula1=f'"{",".join(SOURCES)}"', allow_blank=True)
date_check = DataValidation(type="date", operator="greaterThan", formula1="36526", allow_blank=True,
                            error="Enter a date, for example 01/10/2026.", errorTitle="Date needed")
for validation in (status_list, source_list, date_check):
    ws.add_data_validation(validation)
status_list.add(f"C2:C{ROWS + 1}")
source_list.add(f"J2:J{ROWS + 1}")
for column in ("D", "E", "G"):
    date_check.add(f"{column}2:{column}{ROWS + 1}")
    for row in range(2, ROWS + 2):
        ws[f"{column}{row}"].number_format = "DD/MM/YYYY"

# Highlight overdue follow-ups on active applications, and interviews and offers.
active = 'OR($C2="Saved",$C2="Applied",$C2="Interview",$C2="Offer")'
ws.conditional_formatting.add(f"A2:N{ROWS + 1}", FormulaRule(
    formula=[f'AND({active},$G2<>"",$G2<TODAY())'], fill=PatternFill("solid", fgColor="FFE3E3")))
ws.conditional_formatting.add(f"C2:C{ROWS + 1}", FormulaRule(
    formula=['OR($C2="Interview",$C2="Offer")'], fill=PatternFill("solid", fgColor="F5E6C0")))

guide = wb.create_sheet("How to use")
guide.column_dimensions["A"].width = 100
lines = [
    ("Job application tracker (UK) – free from WorkCV", True),
    ("", False),
    ("1. Add one row for each job you apply for. Job title and Employer are the only must-haves.", False),
    ("2. Pick a Status from the list. Rows with a Next action date in the past turn red until you act.", False),
    ("3. Paste the job advert text so you can tailor your CV and cover letter to it.", False),
    ("4. Review the sheet once a week: chase overdue follow-ups and close finished applications.", False),
    ("", False),
    ("Prefer an online version? Use the free tracker at https://workcv.co.uk/tools/job-application-tracker-uk", False),
    ("To move this sheet there: File > Save As > CSV, then use Import CSV in the online tracker.", False),
    ("", False),
    ("Tailor your CV for each job: https://workcv.co.uk/tools/job-application-pack-uk", False),
]
for row, (text, bold) in enumerate(lines, start=1):
    cell = guide.cell(row=row, column=1, value=text)
    cell.font = Font(bold=bold, size=14 if bold else 11, color=NAVY if bold else "1A1A1A")

wb.active = 0
OUT.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUT)
print(f"wrote {OUT.relative_to(ROOT)} ({OUT.stat().st_size} bytes)")
